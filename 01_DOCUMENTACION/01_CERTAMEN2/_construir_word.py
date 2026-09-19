"""Construye Certamen2_Reconstruccion_Completa.docx a partir de los tres documentos
markdown del expediente. No edita los .md: solo los lee y los compone.

    python _construir_word.py
"""
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt, RGBColor, Cm

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Certamen2_Reconstruccion_Completa.docx")

PARTES = [
    ("Parte I", "Analisis de reconstruccion del proceso de estudio",
     "01_ANALISIS_RECONSTRUCCION_CERTAMEN.md",
     "Analisis completo en cinco fases: identificacion del material, ficha de cada una de "
     "las once preguntas, correspondencias conceptuales, separacion entre conocimiento y "
     "redaccion, y revision adversarial."),
    ("Parte II", "Mapa del proceso de estudio",
     "02_MAPA_PROCESO_ESTUDIO.md",
     "Matriz transversal de las once preguntas y analisis de los patrones globales del "
     "metodo de trabajo documentado en el repositorio."),
    ("Parte III", "Guia de defensa oral",
     "03_DEFENSA_ORAL_CERTAMEN.md",
     "Guion de trabajo por pregunta: concepto central, respuesta corta, razonamiento breve, "
     "ejemplo propio, error a evitar y repregunta probable con su respuesta."),
]

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")


def add_runs(par, text):
    """Escribe texto con negrita, cursiva y codigo inline."""
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            par.add_run(tok[1:-1]).italic = True
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = par.add_run(tok[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x8B, 0x25, 0x00)
        else:
            par.add_run(tok)


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def is_sep(line):
    return bool(re.fullmatch(r"\|[\s:\-\|]+\|", line.strip()))


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        add_runs(cell.paragraphs[0], c)
        for r in cell.paragraphs[0].runs:
            r.bold = True
            r.font.size = Pt(9)
    for row in body:
        cells = t.add_row().cells
        for i, c in enumerate(row[: len(header)]):
            cells[i].text = ""
            add_runs(cells[i].paragraphs[0], c)
            for r in cells[i].paragraphs[0].runs:
                r.font.size = Pt(9)
    doc.add_paragraph()


def render(doc, md, base_level=1):
    """Convierte markdown a contenido de Word. base_level desplaza los encabezados."""
    lines = md.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()

        if not s:
            i += 1
            continue

        # tabla
        if s.startswith("|") and i + 1 < len(lines) and is_sep(lines[i + 1]):
            rows = [split_row(s)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        # bloque de codigo
        if s.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            r = p.add_run("\n".join(buf))
            r.font.name = "Consolas"
            r.font.size = Pt(8.5)
            continue

        # regla horizontal
        if re.fullmatch(r"-{3,}", s):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(6)
            i += 1
            continue

        # encabezados
        m = re.match(r"^(#{1,6})\s+(.*)", s)
        if m:
            lvl = min(len(m.group(1)) + base_level - 1, 6)
            h = doc.add_heading(level=lvl)
            add_runs(h, m.group(2))
            i += 1
            continue

        # cita
        if s.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            add_runs(p, " ".join(buf))
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
            continue

        # vinieta
        if re.match(r"^[-*]\s+", s):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, re.sub(r"^[-*]\s+", "", s))
            i += 1
            continue

        # numerada
        if re.match(r"^\d+\.\s+", s):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\d+\.\s+", "", s))
            i += 1
            continue

        # parrafo
        p = doc.add_paragraph()
        add_runs(p, s)
        i += 1


def main():
    doc = Document()

    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(6)

    # --- Portada ---
    t = doc.add_heading("Certamen 2 - Fundamentos de Ciencias de Datos", level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for txt, sz, bold in [
        ("Reconstruccion del proceso de estudio", 16, True),
        ("Expediente completo", 12, False),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(txt)
        r.font.size = Pt(sz)
        r.bold = bold

    doc.add_paragraph()
    for txt in [
        "Curso: 4321001-0 Fundamentos de Ciencias de Datos (T2-2026)",
        "Universidad de Concepcion",
        "Certamen rendido el 28 de agosto de 2026",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(txt).font.size = Pt(10)

    doc.add_paragraph()
    doc.add_paragraph()

    h = doc.add_heading("Que es este documento", level=1)
    doc.add_paragraph(
        "Un analisis de si el material local del repositorio F:\\MACI permite reconstruir un "
        "proceso de comprension que explique como se resolvieron las once preguntas del "
        "Certamen 2."
    )

    doc.add_heading("Que no es", level=1)
    for txt in [
        "No es una defensa legal ni administrativa. No intenta demostrar inocencia.",
        "No intenta probar que no se uso IA. El uso de IA esta declarado por el autor y no es lo que este documento evalua.",
        "No determina si hubo o no una infraccion academica. Esa es una decision institucional, no una conclusion derivable de archivos.",
        "No usa Git, logs ni timestamps como argumento. Las fechas se mencionan una sola vez, en la revision adversarial, y solo para acotar hasta donde llega la evidencia.",
    ]:
        doc.add_paragraph(txt, style="List Bullet")

    doc.add_heading("La pregunta que si responde", level=1)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.0)
    r = p.add_run(
        "\u00bfPuede reconstruirse, a partir del material local, un proceso de comprension "
        "que explique como se resolvieron las preguntas del certamen?"
    )
    r.italic = True

    doc.add_paragraph(
        "Las tres respuestas admisibles eran: si, parcialmente, o no hay evidencia suficiente."
    )
    p = doc.add_paragraph()
    r = p.add_run("La respuesta obtenida fue: PARCIALMENTE.")
    r.bold = True

    doc.add_heading("Resultado", level=1)
    add_table(doc, [
        ["Nivel de evidencia", "Preguntas", "Total"],
        ["**FUERTE**", "P1 ensambles - P2 sobreajuste - P3 causas - P4 metricas - P7 polinomial - P11 supuestos", "6"],
        ["**PARCIAL**", "P5 LLMs - P6 revoluciones IA - P8 matriz de confusion - P9 curva ROC", "4"],
        ["**INSUFICIENTE**", "-", "0"],
        ["No evaluable", "P10 punto base", "1"],
    ])

    doc.add_heading("Contenido", level=1)
    add_table(doc, [
        ["Parte", "Documento", "Contenido"],
    ] + [[n, t_, d] for n, t_, _f, d in PARTES])

    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run(
        "Regla que rige todo el expediente: no fabricar evidencia, no atribuir pensamientos "
        "no documentados, no modificar archivos originales, y no confundir similitud textual "
        "con comprension. Donde el material no alcanza, el documento lo dice."
    )
    r.italic = True
    r.font.size = Pt(9.5)

    # --- Partes ---
    for nombre, titulo, archivo, _desc in PARTES:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(nombre.upper())
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

        with open(os.path.join(HERE, archivo), encoding="utf-8") as fh:
            md = fh.read()
        render(doc, md, base_level=1)

    doc.save(OUT)
    print("OK ->", OUT)
    print("Parrafos:", len(doc.paragraphs), "| Tablas:", len(doc.tables))


if __name__ == "__main__":
    main()
